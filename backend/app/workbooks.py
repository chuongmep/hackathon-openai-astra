import io
import zipfile
from collections import Counter

import openpyxl

from .config import AppError
from .schemas import Schedule


def open_book(content):
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > 200 * 1024 * 1024:
                raise ValueError("Workbook expands beyond 200 MB")
        return openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    except Exception as exc:
        raise AppError(422, "invalid_workbook", "Upload a valid .xlsx workbook under 200 MB uncompressed") from exc


def cell_value(value):
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    return str(value)


class WorkbookService:
    def __init__(self, storage):
        self.storage = storage

    def upload(self, filename, content):
        if not filename.lower().endswith(".xlsx"):
            raise AppError(422, "invalid_workbook", "Upload an .xlsx workbook")
        book = open_book(content)
        try:
            previews = [{"name": s.title, "rows": s.max_row, "columns": s.max_column,
                         "preview": [[cell_value(c) for c in r] for r in
                                     s.iter_rows(max_row=10, max_col=min(s.max_column or 1, 30), values_only=True)]}
                        for s in book]
        finally:
            book.close()
        return self.storage.public(self.storage.persist_file("workbook", filename, content, {"sheets": previews}))

    def schedule(self, config: Schedule):
        record = self.storage.get("workbook", config.workbook_id)
        book = open_book(self.storage.file(record).read_bytes())
        try:
            if config.sheet not in book.sheetnames:
                raise AppError(422, "invalid_sheet", "Selected worksheet does not exist")
            sheet = book[config.sheet]
            if (sheet.max_row or 0) > 100000 or (sheet.max_column or 0) > 200:
                raise AppError(422, "schedule_too_large", "Schedule is limited to 100,000 rows and 200 columns")
            rows = sheet.iter_rows(min_row=config.header_row, values_only=True)
            headers = [str(x).strip() if x is not None else "" for x in next(rows, ())]
            # Header spelling is case/whitespace insensitive; entity GUIDs remain exact.
            def normalize(value):
                return "".join(value.split()).casefold()
            normalized = [normalize(header) for header in headers]
            columns = [normalize(config.guid_column), normalize(config.material_column)]
            for column, normalized_column in zip((config.guid_column, config.material_column), columns):
                count = normalized.count(normalized_column) if normalized_column else 0
                if count != 1:
                    reason = "Missing" if count == 0 else "Duplicate"
                    available = ", ".join(repr(h) for h in headers if h)[:500] or "(empty row)"
                    raise AppError(422, "invalid_columns", f"{reason} column {column!r} in worksheet {config.sheet!r}, header row {config.header_row}. Available headers: {available}. Choose the correct worksheet, header row and column mappings. Material validation requires IFC GlobalIds and expected materials.")
            if columns[0] == columns[1]:
                raise AppError(422, "invalid_columns", "GUID and material columns must be different")
            gi, mi = (normalized.index(column) for column in columns)
            result = []
            for number, cells in enumerate(rows, start=config.header_row + 1):
                guid, expected = cells[gi], cells[mi]
                if guid is None and expected is None:
                    continue
                result.append({"GlobalId": str(guid) if guid is not None else "",
                               "expected": str(expected) if expected is not None else "", "row": number})
            duplicates = {key for key, count in Counter(r["GlobalId"] for r in result).items() if count > 1}
            return result, duplicates
        finally:
            book.close()
