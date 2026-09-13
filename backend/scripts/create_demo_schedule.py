"""Generate a synthetic schedule tied to actual sample-model GUIDs; never edit the IFC."""
import argparse
from pathlib import Path

import ifcopenshell
import openpyxl
from app.ifc_service import material_data
from openpyxl.styles import Font, PatternFill


def create(source: Path, destination: Path):
    model = ifcopenshell.open(str(source))
    doors = sorted(model.by_type("IfcDoor"), key=lambda d: d.GlobalId)
    if len(doors) < 4:
        raise ValueError("Demo requires at least four doors")
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Door Schedule"
    sheet.append(["GlobalId", "ExpectedMaterial", "DoorName", "DemoNote"])
    for i, door in enumerate(doors[:-1]):
        materials = material_data(door)["names"]
        expected = materials[0] if materials else "DEMO REQUIRED MATERIAL"
        note = "Synthetic expectation copied from model; not a real specification"
        if i == 0:
            expected, note = "DEMO INTENTIONAL MISMATCH", "Intentional mismatch"
        elif i == 1:
            expected, note = None, "Intentional missing expected material"
        elif not materials:
            note = "IFC material missing: expected unknown result"
        sheet.append([door.GlobalId, expected, door.Name, note])
    sheet.append(["DEMO-NOT-IN-MODEL", "Steel", None, "Intentional unresolved GUID"])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="216C62")
    for col, width in {"A": 28, "B": 40, "C": 65, "D": 65}.items():
        sheet.column_dimensions[col].width = width
    notes = book.create_sheet("Demo Instructions")
    notes.append(["SYNTHETIC DEMO — NOT A PROJECT SPECIFICATION"])
    notes.append(["Source IFC", source.name])
    notes.append(["Matching", "Exact GlobalId; first row contains headers"])
    notes.append(["Missing coverage", "Last door intentionally omitted from schedule"])
    notes.append(["Purpose", "Demonstrate pass/fail/unknown and uncovered entities"])
    notes.column_dimensions["A"].width = 60
    notes.column_dimensions["B"].width = 75
    destination.parent.mkdir(parents=True, exist_ok=True)
    book.save(destination)
    return {"door_count": len(doors), "schedule": str(destination), "uncovered_guid": doors[-1].GlobalId}


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    print(create(args.source, args.destination))
